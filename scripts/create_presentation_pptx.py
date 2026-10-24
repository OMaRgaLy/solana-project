import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(7, 11, 20)       # #070B14
    CARD_BG = RGBColor(17, 24, 39)       # #111827
    CARD_BORDER = RGBColor(31, 41, 55)   # #1F2937
    EMERALD = RGBColor(20, 241, 149)     # #14F195
    PURPLE = RGBColor(124, 58, 237)      # #7C3AED
    CYAN = RGBColor(0, 229, 255)         # #00E5FF
    AMBER = RGBColor(245, 158, 11)       # #F59E0B
    WHITE = RGBColor(243, 244, 246)      # #F3F4F6
    MUTED = RGBColor(156, 163, 175)      # #9CA3AF
    RED_ACCENT = RGBColor(239, 68, 68)

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_badge(slide, left, top, text, color=EMERALD, bg_col=RGBColor(15, 30, 28)):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.2), Inches(0.42))
        badge.fill.solid()
        badge.fill.fore_color.rgb = bg_col
        badge.line.color.rgb = color
        badge.line.width = Pt(1)
        tf = badge.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Arial"
        return badge

    def add_header(slide, tag_text, title_text, subtitle_text, tag_color=EMERALD):
        # Tag
        tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(0.4))
        tf = tag.text_frame
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = tag_color
        p.font.name = "Arial"

        # Title
        t = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.8))
        tf2 = t.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.font.name = "Arial"

        # Subtitle
        s = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.5), Inches(0.5))
        tf3 = s.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = subtitle_text
        p3.font.size = Pt(15)
        p3.font.color.rgb = MUTED
        p3.font.name = "Arial"

    # ==========================
    # SLIDE 1: COVER
    # ==========================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Top badges
    add_badge(s1, Inches(0.8), Inches(0.8), "⚡ SOLANA DEVNET · LIVE CODE", EMERALD, RGBColor(12, 35, 25))
    add_badge(s1, Inches(4.3), Inches(0.8), "SUPERTEAM KAZAKHSTAN TRACK", PURPLE, RGBColor(28, 16, 50))

    # Main Title
    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(1.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "SteppeInvoice"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Arial"

    sub_box = s1.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.5), Inches(0.8))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Инвойс в SOL с автоматическим сплитом на 3 кошелька"
    p_sub.font.size = Pt(26)
    p_sub.font.bold = True
    p_sub.font.color.rgb = CYAN
    p_sub.font.name = "Arial"

    desc_box = s1.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.5), Inches(1.2))
    tf_desc = desc_box.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = "Простой некастодиальный инструмент: отправляй ссылку на оплату, получай SOL за 400 мс, делись долей с напарником смарт-контрактом в 1 транзакции и получай официальный Акт с QR-печатью."
    p_desc.font.size = Pt(17)
    p_desc.font.color.rgb = MUTED
    p_desc.font.name = "Arial"

    # Feature pills
    pills = [
        ("✓ 0% КОМИССИИ СЕРВИСА", EMERALD),
        ("✓ АТОМАРНЫЙ СПЛИТ НА КОНТРАКТЕ", PURPLE),
        ("✓ АКТ ВЫПОЛНЕННЫХ РАБОТ", AMBER),
        ("✓ ПОДДЕРЖКА SOLANA BLINKS", CYAN),
    ]
    for i, (text, col) in enumerate(pills):
        box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 2.95), Inches(5.8), Inches(2.8), Inches(0.65))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = col
        box.line.width = Pt(1.5)
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = col
        p.font.name = "Arial"

    # ==========================
    # SLIDE 2: PROBLEM
    # ==========================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Реальная боль фрилансеров", "В чем проблема совместной работы в крипте?", "Рутина, человеческий фактор и отсутствие бухгалтерских документов")

    probs = [
        ("💸", "Ручные переводы и повторные комиссии", "Клиент платит одному человеку. Тот вручную переводит долю напарнику, теряя время на калькулятор и оплачивая комиссии сети повторно.", RED_ACCENT),
        ("🤝", "Человеческий фактор и недоверие", "Вся сумма заказа аккумулируется у одного лица. Возникают риски задержек, разногласий и необходимость слепого доверия.", AMBER),
        ("🏦", "Сложность накоплений на налоги", "Фрилансеру тяжело дисциплинированно откладывать 10-15% с каждого мелкого заказа в отдельную копилку или на налоги.", PURPLE),
        ("📄", "Оплата есть — документов нет", "Хеш транзакции в Solscan непонятен бухгалтеру. Нет официального первичного документа с суммой и датой для спокойной отчетности.", CYAN)
    ]
    coords = [
        (Inches(0.8), Inches(2.2)), (Inches(6.8), Inches(2.2)),
        (Inches(0.8), Inches(4.6)), (Inches(6.8), Inches(4.6)),
    ]
    for i, (icon, title, desc, acc_col) in enumerate(probs):
        l, t = coords[i]
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.7), Inches(2.1))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = acc_col
        c.line.width = Pt(1.5)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{icon}  {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Arial"

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = MUTED
        p2.font.name = "Arial"
        p2.space_before = Pt(8)

    # ==========================
    # SLIDE 3: HOW IT WORKS
    # ==========================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Простой и понятный процесс", "Как устроен SteppeInvoice", "От создания счета до зачисления средств — 4 понятных шага")

    steps = [
        ("1", "Создание счета", "Указывается название задачи, сумма в SOL, кошелек напарника и проценты сплита."),
        ("2", "Отправка ссылки", "Сервис дает прямую веб-ссылку и ссылку Solana Blink для оплаты прямо в соцсетях."),
        ("3", "Оплата в 1 клик", "Заказчик подключает свой Phantom/Solflare и подтверждает платеж за 400 миллисекунд."),
        ("4", "Атомарный расчет", "Смарт-контракт мгновенно начисляет доли на 3 адреса и выдает готовый Акт.")
    ]
    for i, (num, title, desc) in enumerate(steps):
        l = Inches(0.8 + i * 2.95)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(2.4), Inches(2.75), Inches(4.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1)
        tf = c.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = f"Шаг {num}"
        p_num.font.size = Pt(18)
        p_num.font.bold = True
        p_num.font.color.rgb = EMERALD
        p_num.font.name = "Arial"

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE
        p_title.font.name = "Arial"
        p_title.space_before = Pt(10)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = MUTED
        p_desc.font.name = "Arial"
        p_desc.space_before = Pt(14)

    # ==========================
    # SLIDE 4: AUTO-SPLIT
    # ==========================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Ядро технологии", "Атомарный сплит на 3 кошелька", "Смарт-контракт делит входящий платеж в одной транзакции без посредников")

    # Split bar visualization
    bar_bg = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.7), Inches(0.5))
    bar_bg.fill.solid()
    bar_bg.fill.fore_color.rgb = RGBColor(25, 35, 55)
    bar_bg.line.fill.background()

    # Segments
    seg1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.7 * 0.85), Inches(0.5))
    seg1.fill.solid()
    seg1.fill.fore_color.rgb = EMERALD
    seg1.line.fill.background()

    seg2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + 11.7 * 0.85), Inches(2.2), Inches(11.7 * 0.10), Inches(0.5))
    seg2.fill.solid()
    seg2.fill.fore_color.rgb = AMBER
    seg2.line.fill.background()

    seg3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + 11.7 * 0.95), Inches(2.2), Inches(11.7 * 0.05), Inches(0.5))
    seg3.fill.solid()
    seg3.fill.fore_color.rgb = CYAN
    seg3.line.fill.background()

    # 3 Cards below
    splits = [
        ("1. Основной исполнитель", "85% от суммы", "Получает львиную долю оплаты на свой личный кошелек. Все округления (лампорты) всегда остаются у автора инвойса.", EMERALD),
        ("2. Кошелек-копилка (Vault)", "10% от суммы", "Автоматически откладывает заданный процент (на налоги или в накопления) на отдельный безопасный кошелек.", AMBER),
        ("3. Коллаборатор", "5% от суммы", "Напарник мгновенно получает свою оговоренную долю за дизайн или код без задержек и ручных переводов.", CYAN)
    ]
    for i, (title, pct, desc, col) in enumerate(splits):
        l = Inches(0.8 + i * 3.95)
        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(3.0), Inches(3.75), Inches(3.7))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)
        tf = c.text_frame
        tf.word_wrap = True

        p_pct = tf.paragraphs[0]
        p_pct.text = pct
        p_pct.font.size = Pt(22)
        p_pct.font.bold = True
        p_pct.font.color.rgb = col
        p_pct.font.name = "Arial"

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_t.font.name = "Arial"
        p_t.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = MUTED
        p_d.font.name = "Arial"
        p_d.space_before = Pt(12)

    # ==========================
    # SLIDE 5: ACT OF WORKS
    # ==========================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Юридическая опора", "Акт выполненных работ с QR-печатью", "Первичный документ для отчетности, генерируемый в один клик")

    # Left points (2 cards)
    c1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(6.0), Inches(2.1))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📋 Официальные реквизиты и расчет курса"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Arial"
    p2 = tf.add_paragraph()
    p2.text = "В документе фиксируются Исполнитель, Заказчик, предмет работ, сумма в SOL и расчетный эквивалент в тенге (KZT) по официальному курсу."
    p2.font.size = Pt(13)
    p2.font.color.rgb = MUTED
    p2.font.name = "Arial"
    p2.space_before = Pt(8)

    c2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.6), Inches(6.0), Inches(2.1))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🔍 Ончейн-печать (Solscan QR)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Arial"
    p2 = tf2.add_paragraph()
    p2.text = "QR-код на бланке ведет напрямую на транзакцию в Solscan Explorer. Любой бухгалтер или аудитор может проверить факт и секунду закрытия сделки."
    p2.font.size = Pt(13)
    p2.font.color.rgb = MUTED
    p2.font.name = "Arial"
    p2.space_before = Pt(8)

    # Right white document mockup
    doc = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(2.2), Inches(5.3), Inches(4.5))
    doc.fill.solid()
    doc.fill.fore_color.rgb = RGBColor(255, 255, 255)
    doc.line.color.rgb = RGBColor(200, 205, 215)
    tf_doc = doc.text_frame
    tf_doc.word_wrap = True

    p = tf_doc.paragraphs[0]
    p.text = "АКТ ВЫПОЛНЕННЫХ РАБОТ № 0084"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 24, 39)
    p.font.name = "Arial"

    items = [
        ("Исполнитель:", "Фрилансер (Solana Devnet)"),
        ("Заказчик:", "Acme Inc. (Зарубежный клиент)"),
        ("Предмет:", "Разработка интерфейса"),
        ("ИТОГО К ОПЛАТЕ:", "1.00 SOL (~75 000 ₸)"),
        ("Статус расчета:", "✓ ОПЛАЧЕНО В СЕТИ SOLANA (Сплит 3/3)"),
        ("Цифровая печать:", "QR Solscan Transaction Sealed")
    ]
    for label, val in items:
        p = tf_doc.add_paragraph()
        p.text = f"{label} {val}"
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(55, 65, 81)
        p.font.name = "Arial"
        p.space_before = Pt(6)

    # ==========================
    # SLIDE 6: WHAT IS DONE
    # ==========================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Реализованный функционал", "Что готово в проекте сегодня?", "Полностью рабочий продукт на Solana Devnet без сторонних зависимостей")

    feats = [
        ("🦀", "Смарт-контракт на Rust (Anchor)", "Инструкции создания инвойса, атомарной оплаты со сплитом и безопасной отмены инвойса создателем с возвратом ренты.", EMERALD),
        ("⚡", "Поддержка Solana Blinks & Actions", "Готовые API-эндпоинты /api/actions/invoice/... для моментальной оплаты инвойса прямо из ленты соцсетей.", PURPLE),
        ("💻", "Веб-приложение на Next.js", "Удобный интерфейс: страница создания счета, карточка отправки ссылки и экран оплаты для заказчика.", CYAN),
        ("🧪", "Встроенный Демо-режим (5 кошельков)", "Страница /demo с готовыми ключами (Фрилансер, Клиенты, Копилка, Напарник). Тест полного цикла за 20 секунд без расширений!", AMBER)
    ]
    for i, (icon, title, desc, acc_col) in enumerate(feats):
        l, t = coords[i]
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.7), Inches(2.1))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = acc_col
        c.line.width = Pt(1.5)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{icon}  {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Arial"

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = MUTED
        p2.font.name = "Arial"
        p2.space_before = Pt(8)

    output_path = r"c:\Users\hacke\Projects\solana-project\docs\SteppeInvoice_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
